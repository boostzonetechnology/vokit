from __future__ import annotations

import csv
import io
import json
from html.parser import HTMLParser

from control_plane.agents.domain.types import (
    KNOWLEDGE_FILE_EXTENSIONS,
    MAX_KNOWLEDGE_FILE_BYTES,
    MAX_KNOWLEDGE_TEXT_CHARS,
)
from shared_kernel.errors import DomainError


class _HtmlText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = (data or "").strip()
        if text:
            self._parts.append(text)

    def text(self) -> str:
        return "\n".join(self._parts)


def extract_knowledge_text(
    *, filename: str = "", content_type: str = "", data: bytes
) -> str:
    raw = data or b""
    if len(raw) > MAX_KNOWLEDGE_FILE_BYTES:
        raise DomainError("validation_error", "Knowledge file is too large.", http_status=413)
    if not raw:
        raise DomainError("validation_error", "Knowledge file is empty.")
    extension = _extension(filename, content_type)
    if extension not in KNOWLEDGE_FILE_EXTENSIONS:
        raise DomainError(
            "unsupported_knowledge_type",
            "Supported knowledge files are md, txt, pdf, docx, html, csv and json.",
            http_status=422,
        )
    if extension == "pdf":
        text = _pdf(raw)
    elif extension == "docx":
        text = _docx(raw)
    elif extension in {"html", "htm"}:
        text = _html(raw)
    elif extension == "csv":
        text = _csv(_decode(raw))
    elif extension == "json":
        text = _json(_decode(raw))
    else:
        text = _decode(raw)
    cleaned = (text or "").strip()
    if not cleaned:
        raise DomainError(
            "knowledge_extract_failed",
            "No text could be extracted from the file.",
            http_status=422,
        )
    if len(cleaned) > MAX_KNOWLEDGE_TEXT_CHARS:
        cleaned = cleaned[:MAX_KNOWLEDGE_TEXT_CHARS]
    return cleaned


def _extension(filename: str, content_type: str) -> str:
    name = (filename or "").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    if "." in name:
        return name.rsplit(".", 1)[-1].strip().lower()
    mime = (content_type or "").split(";", 1)[0].strip().lower()
    mapping = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "text/markdown": "md",
        "text/x-markdown": "md",
        "text/plain": "txt",
        "text/html": "html",
        "text/csv": "csv",
        "application/json": "json",
    }
    return mapping.get(mime, "")


def _decode(raw: bytes) -> str:
    return raw.decode("utf-8", errors="replace")


def _pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def _docx(raw: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(raw))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _html(raw: bytes) -> str:
    parser = _HtmlText()
    parser.feed(_decode(raw))
    return parser.text()


def _csv(text: str) -> str:
    reader = csv.reader(io.StringIO(text))
    return "\n".join(", ".join(cell.strip() for cell in row if cell.strip()) for row in reader)


def _json(text: str) -> str:
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return text
    return json.dumps(parsed, ensure_ascii=False, indent=2)
